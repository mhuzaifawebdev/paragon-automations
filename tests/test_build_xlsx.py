"""Offline tests for scraper2/build_xlsx.py's Partners-sheet reorder and international-first sort.
Run: python tests/test_build_xlsx.py
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scraper2"))
import build_xlsx as bx  # noqa: E402

results = []


def check(name, got, want):
    ok = got == want
    results.append(ok)
    print(("ok   " if ok else "FAIL ") + name + ("" if ok else f": got {got!r}, want {want!r}"))


header = ["row_id", "institution", "partner_name", "partner_country", "partner_type",
          "mobility_type", "source_url", "evidence_quote", "confidence", "partner_scope",
          "verified", "verify_note", "scraped_at"]
rows = [
    ["r001", "A College", "National Uni", "Ireland", "university", "study",
     "https://a.edu/p", "quote", "0.9", "National", "yes", "", "2026-10-01"],
    ["r001", "A College", "Foreign Uni", "Portugal", "university", "study",
     "https://a.edu/p", "quote", "0.9", "International", "yes", "", "2026-10-01"],
    ["r001", "A College", "EAEC", "", "network", "unknown",
     "https://a.edu/p", "quote", "0.8", "International", "yes", "", "2026-10-01"],
    ["r002", "B College", "Another National", "Spain", "university", "study",
     "https://b.edu/p", "quote", "0.9", "National", "yes", "", "2026-10-01"],
    ["r002", "B College", "Another Foreign", "France", "university", "study",
     "https://b.edu/p", "quote", "0.9", "International", "yes", "", "2026-10-01"],
]

new_header, new_rows = bx._reorder_partners(header, rows)
check("partner_name comes before source_url after reorder",
      new_header.index("partner_name") < new_header.index("source_url"), True)
check("partner_scope comes before source_url after reorder",
      new_header.index("partner_scope") < new_header.index("source_url"), True)
check("reorder drops no columns", set(new_header), set(header))
check("reorder doesn't change row count", len(new_rows), len(rows))

sorted_rows = bx._sort_partners(new_header, new_rows)
name_i, rid_i, scope_i = new_header.index("partner_name"), new_header.index("row_id"), new_header.index("partner_scope")
r001_names = [r[name_i] for r in sorted_rows if r[rid_i] == "r001"]
r002_names = [r[name_i] for r in sorted_rows if r[rid_i] == "r002"]
check("r001: International partners (Foreign Uni, EAEC) come before National (National Uni)",
      r001_names, ["Foreign Uni", "EAEC", "National Uni"])
check("r002: International partner comes before National",
      r002_names, ["Another Foreign", "Another National"])
check("each institute's block stays contiguous (not interleaved with the other institute's rows)",
      [r[rid_i] for r in sorted_rows], ["r001", "r001", "r001", "r002", "r002"])

# Stability: two International partners for the same institute keep their original relative order.
stable_header = ["row_id", "partner_name", "partner_scope"]
stable_rows = [
    ["r001", "First Found", "International"],
    ["r001", "Second Found", "International"],
]
stable_sorted = bx._sort_partners(stable_header, stable_rows)
check("equal-scope rows keep their original (scraped) relative order",
      [r[1] for r in stable_sorted], ["First Found", "Second Found"])

inst_header = ["row_id", "institution_scraped", "office_email", "person_scraped", "email_scraped", "notes"]
inst_new, _ = bx._reorder_institutions(inst_header, [])
check("department email sits directly after the personal email",
      inst_new.index("office_email"), inst_new.index("email_scraped") + 1)
check("institutions reorder drops no columns", set(inst_new), set(inst_header))

# ---- one sheet per institute: only its own partners and campuses ----
import csv  # noqa: E402
import tempfile  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

tmp = Path(tempfile.mkdtemp())


def write(name, hdr, data):
    with open(tmp / name, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(hdr)
        w.writerows(data)


write("institutions_enriched.csv", ["row_id", "institution_scraped", "institution_original", "partners_found", "campuses_found"],
      [["r001", "A College: Arts/Design [Main]", "A", "3", "1"], ["r002", "B College", "B", "2", ""], ["r003", "No Network School", "C", "", ""]])
write("partners.csv", header, rows)
write("campuses.csv", ["row_id", "institution", "campus_name", "country", "city", "role", "source_url", "evidence_quote", "verified", "verify_note", "scraped_at"],
      [["r001", "A College", "North Campus", "Ireland", "Cork", "branch", "https://a.edu/c", "q", "yes", "", ""]])
write("contacts.csv", ["row_id", "name"], [])
write("needs_review.csv", ["row_id", "item"], [])
wb = load_workbook(bx.build(tmp))
own = [s for s in wb.sheetnames if s not in bx.SHEETS]
check("one extra sheet per institute that has partners or campuses, none for the one without", len(own), 2)
check("sheet names are valid for Excel (no forbidden characters, at most 31 long)",
      all(len(s) <= 31 and not any(ch in s for ch in "[]:*?/\\'") for s in own), True)
ws1 = wb[own[0]]
cells = [c for r in ws1.iter_rows(values_only=True) for c in r if c]
check("an institute's sheet lists its own international partners", all(p in cells for p in ("Foreign Uni", "EAEC")), True)
check("partners in the institute's own country are left out", "National Uni" in cells, False)
check("and none of another institute's", any(p in cells for p in ("Another Foreign", "Another National")), False)
check("its campuses are on the same sheet", "North Campus" in cells, True)
check("it links back to the institute's row on the main sheet", ws1["A2"].hyperlink.location, "'Institutions'!A2")
inst_ws = wb["Institutions"]
hdr = [c.value for c in inst_ws[1]]
pf = inst_ws.cell(row=2, column=hdr.index("International Partners Found") + 1)
check("the count on the main sheet is recounted without national partners", pf.value, "2")
check("the partner count on the main sheet opens that institute's own sheet", pf.hyperlink.location, f"'{own[0]}'!A1")
check("the full Partners list has no national partners either", wb["Partners"].max_row, 1 + sum(1 for r in rows if r[9] != "National"))
unk_h = ["row_id", "partner_name", "partner_scope"]
unk = {"Partners": {"header": unk_h, "rows": [["r1", "X", "Unknown"], ["r1", "Y", "National"]]}, "Institutions": {"header": [], "rows": []}}
bx._drop_national(unk)
check("a partner with no country stated is kept, and labelled so", unk["Partners"]["rows"], [["r1", "X", "Country not stated"]])

print(f"\n{sum(results)}/{len(results)} passed")
sys.exit(0 if all(results) else 1)
